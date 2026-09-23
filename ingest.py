"""docs/ 配下の文書を読み込み、チャンク分割してベクトル化し、Chromaへ永続化する。

対応形式: .md / .txt / .pdf / .docx / .xlsx

【このファイルを分けている理由】
Streamlit は画面を操作するたびにスクリプト全体を上から再実行する。
そのため読み込みと Embedding を app.py 側に置くと、質問するたびに
OpenAI の Embedding API を叩き直すことになり、料金も待ち時間も毎回かかる。
「1回だけ走る処理」をこのファイルに閉じ込め、app.py は出来上がった
Chroma を開くだけにする。

【使い方】
    python ingest.py

docs/ の中身を変えたときだけ再実行すればよい。
"""

import os
import sys
from pathlib import Path

from dotenv import load_dotenv

import chromadb
import docx as python_docx
import openpyxl
import pdfplumber
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_openai import OpenAIEmbeddings
from langchain_chroma import Chroma

BASE_DIR = Path(__file__).parent
DOCS_DIR = BASE_DIR / "docs"
PERSIST_DIR = BASE_DIR / "chroma_db"
COLLECTION_NAME = "shanai_docs"

# 同じ内容が複数の形式で置かれている場合に、どれを優先して読むか。
# 実際の共有フォルダにも「報告書.docx」と「報告書.pdf」が並んでいることが多い。
# ファイル名(拡張子を除く)が同じものは、この順で先に来たものだけを索引に入れる。
EXTENSION_PRIORITY = [".docx", ".xlsx", ".pdf", ".md", ".txt"]

# 索引に入れないファイル(サンプルセットの説明書きなど)
EXCLUDE_FILENAMES = {"README.md"}

CHUNK_SIZE = 500
CHUNK_OVERLAP = 50


def load_env_or_exit() -> None:
    load_dotenv(BASE_DIR / ".env")
    if not os.getenv("OPENAI_API_KEY"):
        print(
            "[ERROR] .env に OPENAI_API_KEY が設定されていません。\n"
            "        .env.example をコピーして .env を作り、実際のキーを入れてください。",
            file=sys.stderr,
        )
        sys.exit(1)


# =============================================================================
# 形式ごとの読み込み
#
# どの関数も (テキスト, 位置の説明) の組を並べて返す。
# 「位置の説明」は出典表示に使う。PDFならページ番号、Excelならシート名。
# これがあると商談で「この回答は○○の3ページ目から取りました」と言える。
# =============================================================================


def _read_text_file(path: Path) -> list[tuple[str, str]]:
    # encoding を明示しないと Windows では cp932 で開こうとして
    # UnicodeDecodeError になる。ここは必ず指定する。
    return [(path.read_text(encoding="utf-8"), "")]


def _read_pdf(path: Path) -> list[tuple[str, str]]:
    """ページごとに分けて返す。出典にページ番号を出せるようにするため。"""
    out = []
    with pdfplumber.open(str(path)) as pdf:
        for i, page in enumerate(pdf.pages, start=1):
            text = page.extract_text() or ""
            if text.strip():
                out.append((text, f"{i}ページ"))
    return out


def _read_docx(path: Path) -> list[tuple[str, str]]:
    """段落と表の両方を拾う。

    python-docx の paragraphs には表の中身が含まれない。
    段落だけ読んで表を取りこぼすのは、Word文書を扱うときの定番の事故。
    """
    doc = python_docx.Document(str(path))
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))

    return [("\n".join(parts), "")]


def _read_xlsx(path: Path) -> list[tuple[str, str]]:
    """シートごとに、各行を「列名: 値」の形に展開して返す。

    セルの値をそのまま並べると「A社 15:00 専用パレット」のようになり、
    どの値が何の項目なのかが失われる。そうなると
    「A社の受入締切は?」と聞かれても検索が当たらない。
    見出し行を見つけて「取引先: A社 / 受入締切: 15:00」の形にしておくと、
    行そのものが意味を持った文章になり、検索でも回答でも効く。

    data_only=True が重要。これを付けないと数式のセルが
    "=基準!$B$8*B5" という文字列で返ってきて、計算結果が読めない。
    (ただし一度もExcelで開かれていないファイルは計算結果が保存されて
     いないことがあり、その場合は None になる)
    """
    wb = openpyxl.load_workbook(str(path), data_only=True)
    out = []

    for ws in wb.worksheets:
        rows = [
            [("" if c.value is None else str(c.value).strip()) for c in row]
            for row in ws.iter_rows()
            if any(c.value is not None for c in row)
        ]

        if not rows:
            continue

        # 見出し行 = 2つ以上の値が埋まっている最初の行
        header_idx = None
        for i, row in enumerate(rows):
            if sum(1 for v in row if v) >= 2:
                header_idx = i
                break

        lines = []
        if header_idx is None:
            lines = [" ".join(v for v in row if v) for row in rows]
        else:
            # 見出し行より前(タイトルや注記)はそのまま入れる
            for row in rows[:header_idx]:
                joined = " ".join(v for v in row if v)
                if joined:
                    lines.append(joined)

            headers = rows[header_idx]
            lines.append(" / ".join(v for v in headers if v))

            for row in rows[header_idx + 1:]:
                pairs = [
                    f"{headers[i]}: {v}"
                    for i, v in enumerate(row)
                    if v and i < len(headers) and headers[i]
                ]
                if pairs:
                    lines.append(" / ".join(pairs))

        text = "\n".join(lines)
        if text.strip():
            out.append((text, f"{ws.title} シート"))

    return out


READERS = {
    ".md": _read_text_file,
    ".txt": _read_text_file,
    ".pdf": _read_pdf,
    ".docx": _read_docx,
    ".xlsx": _read_xlsx,
}


# =============================================================================
# タイトルの決定
# =============================================================================


def extract_title(text: str, path: Path) -> str:
    """出典表示に使う文書名を決める。

    ファイル名をそのまま使うと、Windows由来のZIPやNAS経由のマウントで
    文字化けが混ざることがある。本文の先頭にある見出しを優先して使う。
    """
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue
        if line.startswith("# "):
            return line[2:].strip()
        # Markdown以外は先頭の非空行を見出しとみなす。長すぎる場合は本文とみなす。
        if len(line) <= 40:
            return line
        break
    return path.stem


# =============================================================================
# 収集
# =============================================================================


def collect_paths() -> list[Path]:
    """読み込むファイルを決める。同名別形式は優先度の高いものだけ残す。"""
    candidates = [
        p for p in sorted(DOCS_DIR.glob("*"))
        if p.is_file()
        and p.suffix.lower() in READERS
        and p.name not in EXCLUDE_FILENAMES
    ]

    by_stem: dict[str, list[Path]] = {}
    for p in candidates:
        by_stem.setdefault(p.stem, []).append(p)

    chosen, skipped = [], []
    for stem, paths in sorted(by_stem.items()):
        paths.sort(key=lambda p: EXTENSION_PRIORITY.index(p.suffix.lower()))
        chosen.append(paths[0])
        skipped.extend(paths[1:])

    if skipped:
        print("    同名の別形式があるため、次のファイルは読みません:")
        for p in skipped:
            print(f"      - {p.name}")

    return sorted(chosen)


def load_documents() -> list[Document]:
    if not DOCS_DIR.exists():
        print(f"[ERROR] docs フォルダが見つかりません: {DOCS_DIR}", file=sys.stderr)
        sys.exit(1)

    paths = collect_paths()
    if not paths:
        print(f"[ERROR] docs に読み込める文書がありません: {DOCS_DIR}", file=sys.stderr)
        print(f"        対応形式: {', '.join(sorted(READERS))}", file=sys.stderr)
        sys.exit(1)

    documents: list[Document] = []
    for path in paths:
        reader = READERS[path.suffix.lower()]
        try:
            pieces = reader(path)
        except Exception as e:
            print(f"    [警告] 読み込みに失敗したので飛ばします: {path.name} ({e})")
            continue

        if not pieces or not any(t.strip() for t, _ in pieces):
            # スキャンしただけのPDFなど、文字が埋め込まれていないファイルはここに来る。
            # 実案件ではOCRが必要になる合図。
            print(f"    [警告] テキストが取り出せませんでした: {path.name}")
            print("           (画像だけのPDFの可能性があります。OCRが必要です)")
            continue

        title = extract_title(pieces[0][0], path)
        for text, locator in pieces:
            documents.append(Document(
                page_content=text,
                metadata={
                    "source": str(path),
                    "filename": path.name,
                    "filetype": path.suffix.lower().lstrip("."),
                    "title": title,
                    "locator": locator,
                },
            ))

        detail = f"{len(pieces)}区分" if len(pieces) > 1 else ""
        print(f"    読み込み: {path.name} {detail}".rstrip())

    return documents


def split_documents(documents: list[Document]) -> list[Document]:
    """日本語向けに区切り文字を指定して分割する。

    RecursiveCharacterTextSplitter の既定の区切りは英文前提(段落・改行・空白)で、
    日本語だと空白がほとんど無いため、最後の手段である「文字数で強制的に切る」に
    落ちて文の途中で切れやすい。句点と読点を区切り候補に足しておくと、
    チャンクが文の境界で切れるようになり検索精度が上がる。
    """
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=CHUNK_SIZE,
        chunk_overlap=CHUNK_OVERLAP,
        separators=["\n\n", "\n", "。", "、", " ", ""],
    )
    return splitter.split_documents(documents)


def reset_collection() -> None:
    """既存コレクションだけを削除する。

    chroma_db ディレクトリを丸ごと削除する書き方が一般的だが、
    このプロジェクトは OneDrive 配下にあるため同期プロセスがファイルを
    掴んでいて削除に失敗することがある。Chroma 自身のAPIで
    コレクション単位に消すほうが安全。
    """
    client = chromadb.PersistentClient(path=str(PERSIST_DIR))
    try:
        client.delete_collection(COLLECTION_NAME)
        print(f"    既存コレクションを削除しました: {COLLECTION_NAME}")
    except Exception:
        pass  # 初回実行時はコレクションが無いので何もしない


def main() -> None:
    load_env_or_exit()

    print(f"[1] 文書を読み込みます: {DOCS_DIR}")
    documents = load_documents()
    print(f"    → {len(documents)} 件")

    by_type: dict[str, int] = {}
    for d in documents:
        by_type[d.metadata["filetype"]] = by_type.get(d.metadata["filetype"], 0) + 1
    print("    内訳: " + " / ".join(f"{k} {v}" for k, v in sorted(by_type.items())))

    print("[2] チャンクに分割します")
    chunks = split_documents(documents)
    print(f"    → {len(chunks)} チャンク (chunk_size={CHUNK_SIZE}, overlap={CHUNK_OVERLAP})")

    print("[3] 既存インデックスを整理します")
    reset_collection()

    print("[4] ベクトル化して Chroma に保存します(ここで Embedding API を使用)")
    embeddings = OpenAIEmbeddings(model="text-embedding-3-small")
    Chroma.from_documents(
        documents=chunks,
        embedding=embeddings,
        collection_name=COLLECTION_NAME,
        persist_directory=str(PERSIST_DIR),
    )
    print(f"    → 保存先: {PERSIST_DIR}")

    print("\n完了しました。次は  streamlit run app.py  で画面を起動してください。")


if __name__ == "__main__":
    main()
